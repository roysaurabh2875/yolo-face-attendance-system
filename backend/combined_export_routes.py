
from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    LongTable,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.utils import get_column_letter

from backend.database import get_db
from backend.auth_security import require_jwt_admin
from backend.models import (
    AttendanceSession,
    Attendance,
    Subject,
)
from backend.academic_models import AcademicClass
from backend.schedule_models import (
    AttendanceSessionSchedule,
    AttendanceSessionRoster,
)


router = APIRouter(
    prefix="/api/admin/reports/attendance",
    tags=["Admin - Combined Exports"],
    dependencies=[Depends(require_jwt_admin)],
)

MAX_SESSIONS = 500
MAX_RECORDS = 50000


# ==========================================
# 1. SHARED DATA LOADER
# ==========================================

def load_export_data(
    db,
    start_date,
    end_date,
    class_id,
    subject_id,
):
    if start_date > end_date:
        raise HTTPException(
            422, "Start date must not exceed end date"
        )

    if (end_date - start_date).days > 366:
        raise HTTPException(
            422, "Maximum date range is 366 days"
        )

    statement = (
        select(
            AttendanceSession,
            Subject,
            AcademicClass,
        )
        .join(
            AttendanceSessionSchedule,
            AttendanceSessionSchedule.session_id
            == AttendanceSession.session_id,
        )
        .join(
            Subject,
            Subject.subject_id
            == AttendanceSession.subject_id,
        )
        .join(
            AcademicClass,
            AcademicClass.class_id
            == AttendanceSessionSchedule.class_id,
        )
        .where(
            AttendanceSession.status == "FINALIZED",
            AttendanceSession.session_date >= start_date,
            AttendanceSession.session_date <= end_date,
        )
    )

    if class_id is not None:
        statement = statement.where(
            AcademicClass.class_id == class_id
        )

    if subject_id is not None:
        statement = statement.where(
            Subject.subject_id == subject_id
        )

    statement = statement.order_by(
        AttendanceSession.session_date,
        AttendanceSession.period_number,
        AttendanceSession.session_id,
    ).limit(MAX_SESSIONS + 1)

    sessions = db.execute(statement).all()

    if len(sessions) > MAX_SESSIONS:
        raise HTTPException(
            413,
            "Too many sessions. Select a smaller date range.",
        )

    session_ids = [
        session.session_id
        for session, _, _ in sessions
    ]

    if not session_ids:
        raise HTTPException(
            404, "No finalized scheduled sessions found"
        )

    roster_rows = db.scalars(
        select(AttendanceSessionRoster)
        .where(
            AttendanceSessionRoster.session_id.in_(
                session_ids
            )
        )
        .order_by(
            AttendanceSessionRoster.session_id,
            AttendanceSessionRoster.roll_number,
        )
    ).all()

    if len(roster_rows) > MAX_RECORDS:
        raise HTTPException(
            413,
            "Too many student records. Narrow the filters.",
        )

    attendance_rows = db.scalars(
        select(Attendance)
        .where(
            Attendance.session_id.in_(session_ids)
        )
    ).all()

    rosters = {}
    for item in roster_rows:
        rosters.setdefault(item.session_id, []).append(item)

    records = {}
    for item in attendance_rows:
        records.setdefault(item.session_id, {})[
            item.student_id
        ] = item

    details = []
    summaries = []

    overall_present = 0
    overall_absent = 0

    for session, subject, academic_class in sessions:
        roster = rosters.get(session.session_id, [])
        attendance_map = records.get(
            session.session_id, {}
        )

        roster_ids = {
            item.student_id for item in roster
        }

        if not roster or roster_ids != set(attendance_map):
            raise HTTPException(
                409,
                f"Session {session.session_id} "
                "has incomplete or inconsistent attendance",
            )

        present = 0
        absent = 0

        for student in roster:
            attendance = attendance_map[
                student.student_id
            ]

            if attendance.status == "PRESENT":
                present += 1
            elif attendance.status == "ABSENT":
                absent += 1
            else:
                raise HTTPException(
                    409,
                    f"Invalid attendance in session "
                    f"{session.session_id}",
                )

            details.append({
                "date": session.session_date.isoformat(),
                "session_id": session.session_id,
                "period": session.period_number,
                "class_code": academic_class.class_code,
                "subject": subject.subject_code,
                "roll": student.roll_number,
                "name": student.student_name,
                "status": attendance.status,
                "source": attendance.source,
            })

        total = len(roster)

        summaries.append({
            "date": session.session_date.isoformat(),
            "session_id": session.session_id,
            "class_code": academic_class.class_code,
            "subject": subject.subject_code,
            "period": session.period_number,
            "total": total,
            "present": present,
            "absent": absent,
            "percentage": round(
                present / total * 100, 2
            ),
        })

        overall_present += present
        overall_absent += absent

    total_opportunities = (
        overall_present + overall_absent
    )

    return {
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "total_sessions": len(summaries),
        "total_opportunities": total_opportunities,
        "present": overall_present,
        "absent": overall_absent,
        "percentage": round(
            overall_present / total_opportunities * 100,
            2,
        ) if total_opportunities else 0.0,
        "summaries": summaries,
        "details": details,
    }


def safe_excel(value):
    if isinstance(value, str):
        if value.lstrip().startswith(
            ("=", "+", "-", "@", "\t", "\r", "\n")
        ):
            return "'" + value
    return value


def make_filename(start_date, end_date, extension):
    return (
        f"attendance_{start_date.isoformat()}_"
        f"{end_date.isoformat()}.{extension}"
    )


# ==========================================
# 2. COMBINED PDF
# ==========================================

@router.get("/pdf")
def export_combined_pdf(
    start_date: date,
    end_date: date,
    class_id: int | None = Query(default=None, gt=0),
    subject_id: int | None = Query(default=None, gt=0),
    db: Session = Depends(get_db),
):
    report = load_export_data(
        db, start_date, end_date,
        class_id, subject_id,
    )

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph(
            "COMBINED ATTENDANCE REPORT",
            styles["Title"],
        ),
        Spacer(1, 10),
        Paragraph(
            f"<b>Date range:</b> "
            f"{report['start_date']} to "
            f"{report['end_date']}",
            styles["Normal"],
        ),
        Paragraph(
            f"<b>Finalized sessions:</b> "
            f"{report['total_sessions']}",
            styles["Normal"],
        ),
        Paragraph(
            f"<b>Present:</b> {report['present']} "
            f"&nbsp;&nbsp; <b>Absent:</b> "
            f"{report['absent']} "
            f"&nbsp;&nbsp; <b>Attendance:</b> "
            f"{report['percentage']:.2f}%",
            styles["Normal"],
        ),
        Spacer(1, 14),
    ]

    headings = [
        "Date", "Session", "Class", "Subject",
        "Period", "Roll No.", "Student", "Status",
    ]

    table_data = [headings]

    for item in report["details"]:
        values = [
            item["date"],
            str(item["session_id"]),
            item["class_code"],
            item["subject"],
            str(item["period"]),
            item["roll"],
            item["name"],
            item["status"],
        ]

        table_data.append([
            Paragraph(
                escape(str(value)),
                styles["Normal"],
            )
            for value in values
        ])

    table = LongTable(
        table_data,
        colWidths=[
            25*mm, 17*mm, 34*mm, 25*mm,
            16*mm, 39*mm, 79*mm, 30*mm,
        ],
        repeatRows=1,
        hAlign="LEFT",
    )

    table.setStyle(TableStyle([
        (
            "BACKGROUND", (0, 0), (-1, 0),
            colors.HexColor("#203A58")
        ),
        (
            "TEXTCOLOR", (0, 0), (-1, 0),
            colors.white
        ),
        (
            "ROWBACKGROUNDS", (0, 1), (-1, -1),
            [colors.white, colors.HexColor("#F2F5F8")]
        ),
        (
            "GRID", (0, 0), (-1, -1),
            0.3, colors.lightgrey
        ),
        (
            "VALIGN", (0, 0), (-1, -1),
            "MIDDLE"
        ),
        (
            "TOPPADDING", (0, 0), (-1, -1), 6
        ),
        (
            "BOTTOMPADDING", (0, 0), (-1, -1), 6
        ),
    ]))

    # Paragraph header text is styled separately.
    # Set header cells to white as well.
    header_style = styles["Normal"].clone(
        "AttendanceHeader"
    )
    header_style.textColor = colors.white
    header_style.fontName = "Helvetica-Bold"

    table_data[0] = [
        Paragraph(escape(h), header_style)
        for h in headings
    ]

    story.append(table)
    document.build(story)

    buffer.seek(0)

    filename = make_filename(
        start_date, end_date, "pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


# ==========================================
# 3. COMBINED EXCEL
# ==========================================

@router.get("/excel")
def export_combined_excel(
    start_date: date,
    end_date: date,
    class_id: int | None = Query(default=None, gt=0),
    subject_id: int | None = Query(default=None, gt=0),
    db: Session = Depends(get_db),
):
    report = load_export_data(
        db, start_date, end_date,
        class_id, subject_id,
    )

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Session Summary"

    details = workbook.create_sheet(
        "Student Attendance"
    )

    summary.append([
        "Session ID", "Date", "Class", "Subject",
        "Period", "Total", "Present", "Absent",
        "Attendance %",
    ])

    for item in report["summaries"]:
        summary.append([
            item["session_id"],
            item["date"],
            safe_excel(item["class_code"]),
            safe_excel(item["subject"]),
            item["period"],
            item["total"],
            item["present"],
            item["absent"],
            item["percentage"],
        ])

    summary.append([])
    summary.append([
        "OVERALL",
        "", "", "", "",
        report["total_opportunities"],
        report["present"],
        report["absent"],
        report["percentage"],
    ])

    details.append([
        "Date", "Session ID", "Class", "Subject",
        "Period", "Roll Number", "Student Name",
        "Status", "Source",
    ])

    for item in report["details"]:
        details.append([
            item["date"],
            item["session_id"],
            safe_excel(item["class_code"]),
            safe_excel(item["subject"]),
            item["period"],
            safe_excel(item["roll"]),
            safe_excel(item["name"]),
            item["status"],
            safe_excel(item["source"]),
        ])

    for sheet in (summary, details):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (
            f"A1:I{sheet.max_row}"
        )

        for cell in sheet[1]:
            cell.fill = PatternFill(
                "solid", fgColor="203A58"
            )
            cell.font = Font(
                bold=True, color="FFFFFF"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        widths = [
            18, 18, 24, 24, 14,
            22, 28, 18, 20,
        ]

        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[
                get_column_letter(col)
            ].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = make_filename(
        start_date, end_date, "xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )
