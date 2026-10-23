

from datetime import datetime

from io import BytesIO

from xml.sax.saxutils import escape


from fastapi import APIRouter, Depends, HTTPException

from fastapi.responses import StreamingResponse

from sqlalchemy import select

from backend.schedule_models import (

    AttendanceSessionSchedule,

    AttendanceSessionRoster,

)

from sqlalchemy.orm import Session


from reportlab.lib import colors

from reportlab.lib.pagesizes import A4, landscape

from reportlab.lib.styles import getSampleStyleSheet

from reportlab.lib.units import mm

from reportlab.platypus import (

    SimpleDocTemplate, Paragraph, Spacer,

    LongTable, TableStyle

)


from openpyxl import Workbook

from openpyxl.styles import (

    Font, PatternFill, Alignment

)


from backend.database import get_db

from backend.auth_security import get_current_user
from backend.auth_models import UserAccount
from backend.session_permissions import require_session_access

from backend.models import (

    Student, Subject, AttendanceSession, Attendance

)


router = APIRouter(

    prefix="/api/reports",

    tags=["Attendance Reports"]

)


# =====================================

# 1. LOAD FINALIZED ATTENDANCE

# =====================================


def get_report_data(db, session_id):
    """

    Load attendance information for PDF and Excel.



    Scheduled sessions:

        Use frozen roster names and roll numbers.



    Legacy sessions:

        Use existing student records.



    Only FINALIZED sessions can be exported.

    """

    session = db.get(

        AttendanceSession,

        session_id

    )

    if session is None:

        raise HTTPException(

            status_code=404,

            detail="Attendance session not found"

        )

    if session.status != "FINALIZED":

        raise HTTPException(

            status_code=409,

            detail="Finalize attendance before exporting"

        )

    subject = db.get(

        Subject,

        session.subject_id

    )

    if subject is None:

        raise HTTPException(

            status_code=404,

            detail="Subject not found"

        )

    scheduled = db.get(

        AttendanceSessionSchedule,

        session_id

    )

    students = []

    if scheduled is not None:

        # ==================================

        # SCHEDULED SESSION:

        # USE IMMUTABLE ROSTER SNAPSHOT

        # ==================================

        rows = db.execute(

            select(

                AttendanceSessionRoster,

                Attendance

            )

            .outerjoin(

                Attendance,

                (

                    Attendance.session_id

                    == AttendanceSessionRoster.session_id

                )

                &

                (

                    Attendance.student_id

                    == AttendanceSessionRoster.student_id

                )

            )

            .where(

                AttendanceSessionRoster.session_id

                == session_id

            )

            .order_by(

                AttendanceSessionRoster.roll_number

            )

        ).all()

        if not rows:

            raise HTTPException(

                status_code=409,

                detail="Scheduled session roster is empty"

            )

        # Ensure finalized attendance does not

        # contain students outside its roster.

        roster_ids = {

            roster.student_id

            for roster, attendance in rows

        }

        attendance_ids = set(

            db.scalars(

                select(Attendance.student_id)

                .where(

                    Attendance.session_id == session_id

                )

            ).all()

        )

        if attendance_ids != roster_ids:

            raise HTTPException(

                status_code=409,

                detail=(

                    "Finalized attendance does not match "

                    "the frozen session roster"

                )

            )

        for roster, attendance in rows:

            if attendance is None:

                raise HTTPException(

                    status_code=409,

                    detail="Missing finalized attendance record"

                )

            if attendance.status not in (

                "PRESENT", "ABSENT"

            ):

                raise HTTPException(

                    status_code=409,

                    detail="Invalid finalized attendance status"

                )

            students.append({

                "roll": roster.roll_number,

                "name": roster.student_name,

                "status": attendance.status

            })

    else:

        # ==================================

        # LEGACY MANUALLY CREATED SESSIONS

        # ==================================

        rows = db.execute(

            select(Student, Attendance)

            .join(

                Attendance,

                Attendance.student_id

                == Student.student_id

            )

            .where(

                Attendance.session_id == session_id

            )

            .order_by(Student.roll_number)

        ).all()

        if not rows:

            raise HTTPException(

                status_code=409,

                detail="No attendance records found"

            )

        for student, attendance in rows:

            if attendance.status not in (

                "PRESENT", "ABSENT"

            ):

                raise HTTPException(

                    status_code=409,

                    detail="Incomplete attendance records"

                )

            students.append({

                "roll": student.roll_number,

                "name": student.name,

                "status": attendance.status

            })

    present = sum(

        student["status"] == "PRESENT"

        for student in students

    )

    absent = sum(

        student["status"] == "ABSENT"

        for student in students

    )

    total = len(students)

    percentage = round(

        (present / total) * 100,

        2

    ) if total else 0.0

    return {

        "session_id": session_id,

        "subject": subject.subject_name,

        "subject_code": subject.subject_code,

        "section": session.section,

        "date": session.session_date.isoformat(),

        "period": session.period_number,

        "total": total,

        "present": present,

        "absent": absent,

        "percentage": percentage,

        "students": students

    }


# =====================================

# 2. GENERATE PDF

# =====================================


@router.get(

    "/sessions/{session_id}/pdf",


)
def download_pdf(

    session_id: int,

    db: Session = Depends(get_db),
    user: UserAccount = Depends(get_current_user)
):

    require_session_access(db=db, session_id=session_id, user=user)

    report = get_report_data(db, session_id)

    buffer = BytesIO()

    document = SimpleDocTemplate(

        buffer,

        pagesize=landscape(A4),

        rightMargin=18 * mm,

        leftMargin=18 * mm,

        topMargin=15 * mm,

        bottomMargin=15 * mm

    )

    styles = getSampleStyleSheet()

    elements = []

    title = Paragraph(

        "<b>STUDENT ATTENDANCE REPORT</b>",

        styles["Title"]

    )

    elements.append(title)

    elements.append(Spacer(1, 12))

    details = [

        ("Subject", report["subject"]),

        ("Subject Code", report["subject_code"]),

        ("Section", report["section"]),

        ("Date", report["date"]),

        ("Period", str(report["period"]))

    ]

    for label, value in details:

        safe_value = escape(str(value))

        elements.append(

            Paragraph(

                f"<b>{label}:</b> {safe_value}",

                styles["Normal"]

            )

        )

    elements.append(Spacer(1, 16))

    table_data = [[

        "S.No", "Roll Number",

        "Student Name", "Status"

    ]]

    for index, student in enumerate(

        report["students"], start=1

    ):

        table_data.append([

            str(index),

            Paragraph(

                escape(student["roll"]),

                styles["Normal"]

            ),

            Paragraph(

                escape(student["name"]),

                styles["Normal"]

            ),

            student["status"]

        ])

    table = LongTable(

        table_data,

        colWidths=[

            30 * mm,

            55 * mm,

            100 * mm,

            50 * mm

        ],

        repeatRows=1

    )

    table.setStyle(TableStyle([

        (

            "BACKGROUND", (0, 0), (-1, 0),

            colors.HexColor("#203A58")

        ),

        (

            "TEXTCOLOR", (0, 0), (-1, 0),

            colors.white

        ),

        (

            "FONTNAME", (0, 0), (-1, 0),

            "Helvetica-Bold"

        ),

        (

            "ROWBACKGROUNDS", (0, 1), (-1, -1),

            [colors.white, colors.HexColor("#F0F3F6")]

        ),

        (

            "GRID", (0, 0), (-1, -1),

            0.35, colors.lightgrey

        ),

        (

            "VALIGN", (0, 0), (-1, -1),

            "MIDDLE"

        ),

        (

            "BOTTOMPADDING", (0, 0), (-1, -1),

            9

        ),

        (

            "TOPPADDING", (0, 0), (-1, -1),

            9

        )

    ]))

    elements.append(table)

    elements.append(Spacer(1, 18))

    summary = (

        f"<b>Total Students:</b> {report['total']}<br/>"

        f"<b>Present:</b> {report['present']}<br/>"

        f"<b>Absent:</b> {report['absent']}<br/>"

        f"<b>Attendance Percentage:</b> "

        f"{report['percentage']:.2f}%"

    )

    elements.append(

        Paragraph(summary, styles["Normal"])

    )

    elements.append(Spacer(1, 20))

    elements.append(

        Paragraph(

            "Teacher Signature: ____________________",

            styles["Normal"]

        )

    )

    document.build(elements)

    buffer.seek(0)

    filename = (

        f"attendance_session_{session_id}.pdf"

    )

    return StreamingResponse(

        buffer,

        media_type="application/pdf",

        headers={

            "Content-Disposition":

                f'attachment; filename="{filename}"',

            "Cache-Control": "no-store"

        }

    )


# =====================================

# 3. GENERATE EXCEL

# =====================================


@router.get(

    "/sessions/{session_id}/excel",


)
def download_excel(

    session_id: int,

    db: Session = Depends(get_db),
    user: UserAccount = Depends(get_current_user)
):

    require_session_access(db=db, session_id=session_id, user=user)

    report = get_report_data(db, session_id)

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Attendance"

    sheet.merge_cells("A1:D1")

    sheet["A1"] = "STUDENT ATTENDANCE REPORT"

    sheet["A1"].font = Font(

        size=16,

        bold=True,

        color="FFFFFF"

    )

    sheet["A1"].fill = PatternFill(

        "solid",

        fgColor="203A58"

    )

    sheet["A1"].alignment = Alignment(

        horizontal="center"

    )

    def safe_cell(value):

        value = str(value)

        if value.lstrip().startswith(

            ("=", "+", "-", "@")

        ):

            return "'" + value

        return value

    sheet["A3"] = "Subject"

    sheet["B3"] = safe_cell(report["subject"])

    sheet["A4"] = "Subject Code"

    sheet["B4"] = safe_cell(report["subject_code"])

    sheet["A5"] = "Section"

    sheet["B5"] = safe_cell(report["section"])

    sheet["A6"] = "Date"

    sheet["B6"] = report["date"]

    sheet["A7"] = "Period"

    sheet["B7"] = report["period"]

    headers = [

        "S.No",

        "Roll Number",

        "Student Name",

        "Status"

    ]

    header_row = 9

    for column, title in enumerate(headers, start=1):

        cell = sheet.cell(

            row=header_row,

            column=column,

            value=title

        )

        cell.font = Font(

            bold=True,

            color="FFFFFF"

        )

        cell.fill = PatternFill(

            "solid",

            fgColor="203A58"

        )

    for index, student in enumerate(

        report["students"], start=1

    ):

        row = header_row + index

        sheet.cell(row, 1, index)

        sheet.cell(row, 2, safe_cell(student["roll"]))

        sheet.cell(row, 3, safe_cell(student["name"]))

        sheet.cell(row, 4, student["status"])

        status_cell = sheet.cell(row, 4)

        if student["status"] == "PRESENT":

            status_cell.font = Font(

                bold=True,

                color="008000"

            )

        else:

            status_cell.font = Font(

                bold=True,

                color="C00000"

            )

    summary_row = header_row + len(

        report["students"]

    ) + 3

    summaries = [

        ("Total Students", report["total"]),

        ("Present", report["present"]),

        ("Absent", report["absent"]),

        ("Attendance %", report["percentage"])

    ]

    for offset, (label, value) in enumerate(summaries):

        sheet.cell(summary_row + offset, 1, label)

        sheet.cell(summary_row + offset, 2, value)

    sheet.column_dimensions["A"].width = 23

    sheet.column_dimensions["B"].width = 25

    sheet.column_dimensions["C"].width = 36

    sheet.column_dimensions["D"].width = 18

    sheet.freeze_panes = "A10"

    last_row = header_row + len(report["students"])

    sheet.auto_filter.ref = (

        f"A{header_row}:D{last_row}"

    )

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    filename = (

        f"attendance_session_{session_id}.xlsx"

    )

    return StreamingResponse(

        buffer,

        media_type=(

            "application/vnd.openxmlformats-"

            "officedocument.spreadsheetml.sheet"

        ),

        headers={

            "Content-Disposition":

                f'attachment; filename="{filename}"',

            "Cache-Control": "no-store"

        }

    )
